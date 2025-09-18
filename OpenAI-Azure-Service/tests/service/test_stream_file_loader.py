import os
from io import BytesIO
from unittest.mock import MagicMock, patch

import openpyxl
from docx import Document

from src.service.StreamFileLoader import (
    StreamFileLoader,
    _parse_csv,
    _parse_docx,
    _parse_pdf,
    _parse_xlsx,
    _split_content_into_segments,
)


class TestStreamFileLoader:

    @patch("src.service.StreamFileLoader.write_data_in_temporary_file")
    @patch("src.service.StreamFileLoader.fitz.open")
    @patch("os.remove")
    def test_parse_pdf(self, mock_os_remove, mock_fitz_open, mock_write_data):
        mock_write_data.return_value = "temp_file_path"
        mock_doc = MagicMock()
        mock_page = MagicMock()
        mock_page.get_text.return_value = "Sample text"
        mock_doc.page_count = 1
        mock_doc.load_page.return_value = mock_page
        mock_fitz_open.return_value.__enter__.return_value = mock_doc

        data = BytesIO(b"%PDF-1.4\n...")
        result = _parse_pdf(data, "sample.pdf")
        assert "Sample text" in result

    def test_parse_csv(self):
        data = BytesIO(b"col1,col2\nval1,val2\n")
        result = _parse_csv(data, "sample.csv")
        expected_result = "sample.csv:col1,col2\nval1,val2\n".replace("\n", os.linesep)
        assert expected_result in result

    def test_parse_xlsx(self):
        data = BytesIO()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["col1", "col2"])
        ws.append(["val1", "val2"])
        wb.save(data)
        data.seek(0)

        result = _parse_xlsx(data, "sample.xlsx")
        assert "sample.xlsx:col1\tcol2\nval1\tval2" in result

    def test_parse_docx(self):
        data = BytesIO()
        doc = Document()
        doc.add_paragraph("Sample text")
        doc.save(data)
        data.seek(0)

        result = _parse_docx(data, "sample.docx")
        assert "sample.docx:Sample text" in result

    @patch("src.service.StreamFileLoader.write_data_in_temporary_file")
    @patch("src.service.StreamFileLoader.fitz.open")
    @patch("os.remove")
    @patch("src.service.StreamFileLoader.Blob.from_data")
    def test_load_list(
        self, mock_blob_from_data, mock_os_remove, mock_fitz_open, mock_write_data
    ):
        mock_write_data.return_value = "temp_file_path"
        mock_doc = MagicMock()
        mock_page = MagicMock()
        mock_page.get_text.return_value = "Sample text"
        mock_doc.page_count = 1
        mock_doc.load_page.return_value = mock_page
        mock_fitz_open.return_value.__enter__.return_value = mock_doc
        mock_blob_from_data.side_effect = lambda data, path: MagicMock(
            data=data, source=path
        )

        data_list = [BytesIO(b"%CSV-1.4\n..."), BytesIO(b"%PDF-1.4\n...")]
        file_names = ["sample.csv", "sample.pdf"]

        loader = StreamFileLoader(data_list=data_list, file_names=file_names)
        documents = loader.load_list()

        assert len(documents) == 1
        assert any("sample.csv" in doc.page_content for doc in documents)
        assert any("Sample text" in doc.page_content for doc in documents)

    def test_split_content_into_segments_else(self):
        content = "line1\n" + "line2\n" * 100000
        segments = _split_content_into_segments(content)
        assert len(segments) == 3
        assert any("line2" in segment for segment in segments)
